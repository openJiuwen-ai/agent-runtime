from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from conftest import ManagerApiHarness
from demo_payloads import model_templates
from openpyxl import load_workbook

from manager_server.core.authz import AuthzService
from manager_server.core.import_export.registry import (
    ImportExportContext,
    SheetData,
    WorkbookData,
)
from manager_server.core.import_export.standalone import OrganizationAdapter, RoleAdapter
from manager_server.core.import_export.workbook import SENSITIVE_REDACTED
from manager_server.models.authz_models import AUTHZ_ROLE_USER_TABLE_DEF


def _fill_model_api_key(payload: bytes, api_key: str) -> bytes:
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["12_ModelTemplates"]
    headers = [str(cell.value or "") for cell in sheet[1]]
    column = headers.index("api_key") + 1
    assert sheet.cell(2, column).value == SENSITIVE_REDACTED
    sheet.cell(2, column).value = api_key
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.asyncio
async def test_model_xlsx_round_trip_preserves_business_id(manager_api: ManagerApiHarness):
    h = manager_api
    _, create_body = model_templates()[0]
    created = await h.post_json("/model-templates", create_body)
    template_id = created["template_id"]

    exported = await h.http.post(
        "/api/v1/import-export/model/export",
        json={"resource_ids": [template_id]},
    )
    assert exported.status_code == 200, exported.text
    payload = _fill_model_api_key(exported.content, create_body["api_key"])

    await h.delete_ok(f"/model-templates/{template_id}")
    preflight = await h.http.post(
        "/api/v1/import-export/model/preflight",
        content=payload,
        headers={
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        },
    )
    assert preflight.status_code == 200, preflight.text
    report = preflight.json()["data"]
    assert report["can_import"] is True
    assert report["summary"]["create"] == 1

    imported = await h.http.post(
        "/api/v1/import-export/model/import",
        params={"confirmation_token": report["confirmation_token"]},
        content=payload,
        headers={
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        },
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["data"]["created_manager_objects"] == 1

    restored = await h.get_json(f"/model-templates/{template_id}")
    assert restored["template_id"] == template_id
    for key, value in create_body.items():
        assert restored[key] == value

    duplicate = await h.http.post(
        "/api/v1/import-export/model/preflight",
        content=payload,
        headers={
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        },
    )
    assert duplicate.status_code == 200
    assert duplicate.json()["data"]["summary"]["reuse"] == 1


@pytest.mark.asyncio
async def test_organization_export_resolves_member_identity_provider(
    manager_api: ManagerApiHarness,
    monkeypatch: pytest.MonkeyPatch,
):
    async def fake_identity_request(
        method: str,
        path: str,
        _authorization: str | None,
        body: dict | None = None,
    ):
        assert method == "GET"
        assert body is None
        if path == "/v1/orgs/org-1":
            return {"group_id": "org-1", "display_name": "Organization 1", "status": "active"}
        if path == "/v1/orgs/org-1/members":
            # The real members endpoint does not include identity_provider.
            return {"users": [{"user_id": "local-user", "display_name": "Local User"}]}
        if path == "/v1/users/local-user":
            return {
                "user_id": "local-user",
                "username": "local-user",
                "identity_provider": "local",
                "display_name": "Local User",
                "status": "active",
            }
        raise AssertionError(f"unexpected identity request: {path}")

    monkeypatch.setattr(
        "manager_server.core.import_export.standalone._identity_request",
        fake_identity_request,
    )
    monkeypatch.setattr(
        "manager_server.core.import_export.cluster._identity_request",
        fake_identity_request,
    )
    workbook = await OrganizationAdapter().export_many(
        ImportExportContext(manager_api.handler, "Bearer test"),
        ["org-1"],
    )

    users = next(sheet for sheet in workbook.sheets if sheet.name == "02_Users")
    assert users.rows == [
        {
            "user_id": "local-user",
            "username": "local-user",
            "identity_provider": "local",
            "display_name": "Local User",
            "status": "active",
            "initial_password": "",
        }
    ]


@pytest.mark.asyncio
async def test_existing_admin_role_does_not_warn_about_skipped_assignments(
    manager_api: ManagerApiHarness,
):
    h = manager_api
    role = await AuthzService(h.handler).create_role(
        role_id="existing_admin_role",
        name="Existing Admin Role",
        description="Already present",
        scope="admin",
        permission_ids=["quota:read"],
        operator_id="test-operator",
    )
    workbook = WorkbookData(
        resource_type="role",
        resource_id="existing_admin_role",
        resource_name="Role",
        sheets=[
            SheetData(
                "02_Users",
                [
                    "user_id",
                    "username",
                    "identity_provider",
                    "display_name",
                    "status",
                    "initial_password",
                ],
                [
                    {
                        "user_id": "existing-user",
                        "username": "existing-user",
                        "identity_provider": "local",
                        "display_name": "Existing User",
                        "status": "active",
                        "initial_password": "",
                    }
                ],
            ),
            SheetData(
                "40_Roles",
                ["role_id", "name", "description", "scope", "is_system", "enabled"],
                [
                    {
                        "role_id": role["role_id"],
                        "name": role["name"],
                        "description": role["description"],
                        "scope": role["scope"],
                        "is_system": role["is_system"],
                        "enabled": role["enabled"],
                    }
                ],
            ),
            SheetData(
                "41_RolePermissions",
                ["role_id", "permission_id"],
                [{"role_id": role["role_id"], "permission_id": "quota:read"}],
            ),
            SheetData(
                "42_RoleAssignments",
                ["role_id", "user_id", "expires_at"],
                [
                    {
                        "role_id": role["role_id"],
                        "user_id": "existing-user",
                        "expires_at": None,
                    }
                ],
            ),
        ],
    )

    report = await RoleAdapter().preflight(
        ImportExportContext(manager_api.handler, "Bearer test"),
        workbook,
    )

    assert report["can_import"] is True
    assert report["summary"] == {"reuse": 1}


@pytest.mark.asyncio
async def test_organization_import_creates_missing_local_user_and_membership(
    manager_api: ManagerApiHarness,
    monkeypatch: pytest.MonkeyPatch,
):
    users: dict[str, dict] = {}
    organizations: dict[str, dict] = {}
    memberships: dict[str, list[str]] = {}

    async def fake_identity_request(
        method: str,
        path: str,
        _authorization: str | None,
        body: dict | None = None,
    ):
        if method == "GET" and path == "/v1/orgs/new-org":
            return organizations.get("new-org")
        if method == "GET" and path == "/v1/users/new-user":
            return users.get("new-user")
        if method == "GET" and path == "/v1/users/by-username/new-user":
            return None
        if method == "POST" and path == "/v1/orgs/":
            assert body == {"group_id": "new-org", "display_name": "New Organization"}
            organizations["new-org"] = {**body, "status": "active"}
            return organizations["new-org"]
        if method == "POST" and path == "/v1/users/":
            assert body == {
                "user_id": "new-user",
                "username": "new-user",
                "display_name": "New User",
                "is_admin": False,
                "password": "initial-password",
            }
            users["new-user"] = {
                **body,
                "identity_provider": "local",
                "status": "active",
            }
            return users["new-user"]
        if method == "POST" and path == "/v1/orgs/new-org/members":
            assert body == {"user_ids": ["new-user"]}
            memberships["new-org"] = list(body["user_ids"])
            return {"user_ids": memberships["new-org"]}
        raise AssertionError(f"unexpected identity request: {method} {path} {body}")

    monkeypatch.setattr(
        "manager_server.core.import_export.standalone._identity_request",
        fake_identity_request,
    )
    monkeypatch.setattr(
        "manager_server.core.import_export.cluster._identity_request",
        fake_identity_request,
    )
    workbook = WorkbookData(
        resource_type="organization",
        resource_id="new-org",
        resource_name="Organization",
        sheets=[
            SheetData(
                "02_Users",
                [
                    "user_id",
                    "username",
                    "identity_provider",
                    "display_name",
                    "status",
                    "initial_password",
                ],
                [
                    {
                        "user_id": "new-user",
                        "username": "new-user",
                        "identity_provider": "local",
                        "display_name": "New User",
                        "status": "active",
                        "initial_password": "initial-password",
                    }
                ],
            ),
            SheetData(
                "03_Organizations",
                ["group_id", "display_name", "status"],
                [
                    {
                        "group_id": "new-org",
                        "display_name": "New Organization",
                        "status": "active",
                    }
                ],
            ),
            SheetData(
                "22_OrganizationMembers",
                ["group_id", "user_id"],
                [{"group_id": "new-org", "user_id": "new-user"}],
            ),
        ],
    )

    result = await OrganizationAdapter().apply(
        ImportExportContext(manager_api.handler, "Bearer test"),
        workbook,
    )

    assert result["created_identity_objects"] == [
        {"object_type": "organization", "object_id": "new-org"},
        {"object_type": "user", "object_id": "new-user"},
    ]
    assert memberships == {"new-org": ["new-user"]}


@pytest.mark.asyncio
async def test_role_import_creates_assignment_with_expiry(
    manager_api: ManagerApiHarness,
    monkeypatch: pytest.MonkeyPatch,
):
    expiry = datetime(2030, 1, 15, 12, 0, tzinfo=UTC)

    async def fake_identity_request(
        method: str,
        path: str,
        _authorization: str | None,
        body: dict | None = None,
    ):
        assert method == "GET"
        assert path == "/v1/users/existing-user"
        assert body is None
        return {
            "user_id": "existing-user",
            "username": "existing-user",
            "identity_provider": "local",
            "display_name": "Existing User",
            "status": "active",
        }

    monkeypatch.setattr(
        "manager_server.core.import_export.standalone._identity_request",
        fake_identity_request,
    )
    monkeypatch.setattr(
        "manager_server.core.import_export.cluster._identity_request",
        fake_identity_request,
    )
    workbook = WorkbookData(
        resource_type="role",
        resource_id="new_admin_role",
        resource_name="Role",
        sheets=[
            SheetData(
                "02_Users",
                [
                    "user_id",
                    "username",
                    "identity_provider",
                    "display_name",
                    "status",
                    "initial_password",
                ],
                [
                    {
                        "user_id": "existing-user",
                        "username": "existing-user",
                        "identity_provider": "local",
                        "display_name": "Existing User",
                        "status": "active",
                        "initial_password": "",
                    }
                ],
            ),
            SheetData(
                "40_Roles",
                ["role_id", "name", "description", "scope", "is_system", "enabled"],
                [
                    {
                        "role_id": "new_admin_role",
                        "name": "New Admin Role",
                        "description": "Created from XLSX",
                        "scope": "admin",
                        "is_system": False,
                        "enabled": True,
                    }
                ],
            ),
            SheetData(
                "41_RolePermissions",
                ["role_id", "permission_id"],
                [{"role_id": "new_admin_role", "permission_id": "quota:read"}],
            ),
            SheetData(
                "42_RoleAssignments",
                ["role_id", "user_id", "expires_at"],
                [
                    {
                        "role_id": "new_admin_role",
                        "user_id": "existing-user",
                        "expires_at": expiry,
                    }
                ],
            ),
        ],
    )
    adapter = RoleAdapter()

    report = await adapter.preflight(
        ImportExportContext(manager_api.handler, "Bearer test"),
        workbook,
    )
    result = await adapter.apply(
        ImportExportContext(
            manager_api.handler,
            "Bearer test",
            operator_id="test-operator",
        ),
        workbook,
    )

    assert report["summary"] == {"create": 1, "warning": 1, "reuse": 1}
    assert result["created_manager_objects"] == 2
    assignment = await manager_api.handler.get(
        AUTHZ_ROLE_USER_TABLE_DEF.table_name,
        {"role_id": "new_admin_role", "user_id": "existing-user"},
    )
    assert assignment is not None
    assert assignment.expires_at.replace(tzinfo=UTC) == expiry
