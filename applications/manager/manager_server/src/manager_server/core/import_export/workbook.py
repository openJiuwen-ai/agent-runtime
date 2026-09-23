from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from .registry import SheetData, WorkbookData

SENSITIVE_REDACTED = "__SENSITIVE_REDACTED__"
MAX_WORKBOOK_BYTES = 20 * 1024 * 1024

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_ID_FILL = PatternFill("solid", fgColor="E7E6E6")
_SENSITIVE_FILL = PatternFill("solid", fgColor="FFF2CC")
_README_FILL = PatternFill("solid", fgColor="D9EAF7")
_TAB_COLORS = {
    "0": "5B9BD5",
    "1": "70AD47",
    "2": "ED7D31",
    "3": "A5A5A5",
    "9": "8064A2",
}
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool, datetime, date)):
        return value
    return str(value)


def _append_safe_row(ws, values: list[Any]) -> None:
    """Append values without allowing configuration text to become an Excel formula."""
    cell_values = [_cell_value(value) for value in values]
    ws.append(cell_values)
    row_index = ws.max_row
    for column_index, value in enumerate(cell_values, start=1):
        if isinstance(value, str) and value.lstrip().startswith(_FORMULA_PREFIXES):
            ws.cell(row_index, column_index).data_type = "s"


def _autosize(ws) -> None:
    for index, column in enumerate(ws.columns, start=1):
        width = 10
        for cell in list(column)[:500]:
            value = "" if cell.value is None else str(cell.value)
            width = max(
                width, min(55, max((len(line) for line in value.splitlines()), default=0) + 2)
            )
        ws.column_dimensions[get_column_letter(index)].width = width


def _apply_sheet_style(ws, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(len(headers), 1))}{max(ws.max_row, 1)}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _TAB_COLORS.get(ws.title[:1], "5B9BD5")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_index, header in enumerate(headers, start=1):
        if header.endswith("_id") or header in {
            "jiuwenclaw_id",
            "resource_id",
            "template_id",
            "container_id",
            "policy_id",
        }:
            for row_index in range(2, ws.max_row + 1):
                ws.cell(row_index, col_index).fill = _ID_FILL
        if header in {
            "enabled",
            "is_admin",
            "allow_http",
            "allow_private_network",
            "allow_public_http",
        }:
            validation = DataValidation(type="list", formula1='"true,false"', allow_blank=True)
            ws.add_data_validation(validation)
            validation.add(f"{get_column_letter(col_index)}2:{get_column_letter(col_index)}1048576")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(max(len(headers), 1))}{max(ws.max_row, 2)}",
        CellIsRule(operator="equal", formula=[f'"{SENSITIVE_REDACTED}"'], fill=_SENSITIVE_FILL),
    )
    _autosize(ws)


def dump_workbook(data: WorkbookData) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    readme = wb.create_sheet("00_ReadMe")
    readme_headers = [
        "format_version",
        "resource_type",
        "source_resource_id",
        "source_resource_name",
        "exported_at",
        "sheet_name",
        "row_count",
        "notes",
    ]
    readme.append(readme_headers)
    exported_at = datetime.now(UTC).isoformat()
    for sheet in data.sheets:
        _append_safe_row(
            readme,
            [
                data.format_version,
                data.resource_type,
                data.resource_id,
                data.resource_name,
                exported_at,
                sheet.name,
                len(sheet.rows),
                sheet.notes,
            ],
        )
        link_cell = readme.cell(readme.max_row, 6)
        link_cell.hyperlink = f"#{quote_sheetname(sheet.name)}!A1"
        link_cell.style = "Hyperlink"
    _apply_sheet_style(readme, readme_headers)
    for row in readme.iter_rows(min_row=2, max_col=len(readme_headers)):
        for cell in row:
            if cell.column <= 5:
                cell.fill = _README_FILL

    for sheet in data.sheets:
        ws = wb.create_sheet(sheet.name)
        ws.append(sheet.headers)
        for row_index, record in enumerate(sheet.rows, start=2):
            _append_safe_row(ws, [record.get(header) for header in sheet.headers])
            for header in sheet.headers:
                if (row_index - 2, header) in sheet.sensitive_cells:
                    ws.cell(row_index, sheet.headers.index(header) + 1).fill = _SENSITIVE_FILL
        _apply_sheet_style(ws, sheet.headers)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def load_workbook_data(raw: bytes, *, expected_resource_type: str) -> WorkbookData:
    if not raw:
        raise ValueError("workbook is empty")
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise ValueError("workbook exceeds 20 MiB limit")
    try:
        wb = load_workbook(BytesIO(raw), read_only=False, data_only=False)
    except Exception as exc:
        raise ValueError(f"invalid xlsx workbook: {exc}") from exc
    if "00_ReadMe" not in wb.sheetnames:
        raise ValueError("missing required sheet: 00_ReadMe")
    readme = wb["00_ReadMe"]
    headers = [str(cell.value or "").strip() for cell in readme[1]]
    metadata_rows = [
        dict(zip(headers, values, strict=False))
        for values in readme.iter_rows(min_row=2, values_only=True)
    ]
    if not metadata_rows:
        raise ValueError("00_ReadMe contains no sheet manifest")
    first = metadata_rows[0]
    resource_type = str(first.get("resource_type") or "").strip().lower()
    if resource_type != expected_resource_type.lower():
        raise ValueError(
            f"workbook resource_type {resource_type!r} does not match {expected_resource_type!r}"
        )
    format_version = str(first.get("format_version") or "").strip()
    if format_version != "1.0":
        raise ValueError(f"unsupported format_version: {format_version}")

    sheets: list[SheetData] = []
    for manifest in metadata_rows:
        name = str(manifest.get("sheet_name") or "").strip()
        if not name:
            continue
        if name not in wb.sheetnames:
            raise ValueError(f"manifest sheet missing: {name}")
        ws = wb[name]
        sheet_headers = [str(cell.value or "").strip() for cell in ws[1]]
        if not all(sheet_headers):
            raise ValueError(f"blank header in sheet: {name}")
        rows = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if all(value is None or value == "" for value in values):
                continue
            rows.append(dict(zip(sheet_headers, values, strict=False)))
        sheets.append(
            SheetData(
                name=name,
                headers=sheet_headers,
                rows=rows,
                notes=str(manifest.get("notes") or ""),
            )
        )
    return WorkbookData(
        resource_type=resource_type,
        resource_id=str(first.get("source_resource_id") or ""),
        resource_name=str(first.get("source_resource_name") or ""),
        format_version=format_version,
        sheets=sheets,
    )
